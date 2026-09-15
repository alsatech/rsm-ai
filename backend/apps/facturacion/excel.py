from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

MESES_EN = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')

HEADER_FILL = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True)
MONEY_FORMAT = '$#,##0.00'
COLUMN_WIDTHS = {'A': 15, 'B': 45, 'C': 35, 'D': 15}
ENCABEZADOS = ('FECHA', 'FACTURA', 'CONCEPTO', 'IMPORTE')


def _fecha_texto(fecha: date) -> str:
    # Texto fijo en inglés (no depende del locale de Excel al abrirlo) para igualar
    # exactamente el formato que ya usa la contadora, ej. "24-Apr-26".
    return f'{fecha.day:02d}-{MESES_EN[fecha.month - 1]}-{fecha.strftime("%y")}'


def generar_excel_relacion(facturas):
    """facturas: queryset/lista de Factura ordenada por fecha ascendente."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Relación'

    for columna, ancho in COLUMN_WIDTHS.items():
        ws.column_dimensions[columna].width = ancho

    for idx, texto in enumerate(ENCABEZADOS, start=1):
        celda = ws.cell(row=2, column=idx, value=texto)
        celda.fill = HEADER_FILL
        celda.font = HEADER_FONT

    fila = 3
    total = 0
    for factura in facturas:
        ws.cell(row=fila, column=1, value=_fecha_texto(factura.fecha))
        ws.cell(row=fila, column=2, value=factura.numero_factura)
        ws.cell(row=fila, column=3, value=factura.concepto)
        celda_importe = ws.cell(row=fila, column=4, value=float(factura.importe))
        celda_importe.number_format = MONEY_FORMAT
        total += factura.importe
        fila += 1

    celda_suma_label = ws.cell(row=fila, column=3, value='SUMA')
    celda_suma_label.font = Font(bold=True)
    celda_suma_label.alignment = Alignment(horizontal='right')
    celda_suma_valor = ws.cell(row=fila, column=4, value=float(total))
    celda_suma_valor.font = Font(bold=True)
    celda_suma_valor.number_format = MONEY_FORMAT

    return wb, total
